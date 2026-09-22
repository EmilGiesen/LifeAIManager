#!/usr/bin/env python3
"""
read_messages.py — holt neue Nachrichten aus den Logineo-Kursräumen (Matrix)
und gibt sie als EINE JSON-Zeile auf stdout aus (für n8n).

Alles andere (Debug, Warnungen) geht nach stderr.

Neu gegenüber der alten Version:
  1. NUR LESEN: Das Skript schickt keine Nachrichten, keine Lesebestätigungen,
     keine Schlüssel-Anfragen und keine Geräte-Nachrichten an andere. Einzige
     Schreibaktion: einmalig die eigenen öffentlichen Geräteschlüssel
     veröffentlichen (keys_upload), damit der Pi Raum-Schlüssel empfangen kann.
  2. Verschlüsselte Anhänge (RoomEncryptedFile/Image/Video/Audio) werden
     heruntergeladen UND entschlüsselt.
  3. Nicht entschlüsselbare Nachrichten landen in pending_megolm.json und
     werden bei jedem Lauf erneut versucht (z.B. nach einem Schlüssel-Import).
  4. Timeline-Limit 100 statt 10 + Nachladen, falls trotzdem abgeschnitten.
  5. Anzeigename des Absenders ("sender_name") für bessere Zusammenfassungen.
  6. Text aus PDF-/Word-Anhängen ("text_auszug"), damit Claude Arbeitsblätter lesen kann
     (braucht: ./venv/bin/pip install pypdf python-docx).
"""
import asyncio
import json
import logging
import re
import sys
import time
from pathlib import Path

from nio import (
    AsyncClient,
    AsyncClientConfig,
    DownloadError,
    MegolmEvent,
    RoomEncryptedAudio,
    RoomEncryptedFile,
    RoomEncryptedImage,
    RoomEncryptedVideo,
    RoomGetEventError,
    RoomMessageAudio,
    RoomMessageFile,
    RoomMessageImage,
    RoomMessagesError,
    RoomMessageText,
    RoomMessageVideo,
    SyncError,
)
from nio.crypto.attachments import decrypt_attachment

from sync_state import load_token, save_token

# nio meldet sonst hunderte harmlose Warnungen (Schema-Fehler, "undecryptable
# Megolm event" usw.). Unsere eigene Zusammenfassung auf stderr reicht.
logging.getLogger("nio").setLevel(logging.CRITICAL)

HOMESERVER = "https://190688.logineonrw-messenger.de"
USER_ID = "@dqq989d8:190688.logineonrw-messenger.de"
BASE = Path(__file__).parent
# Zugangsdaten stehen NICHT mehr hier, sondern in matrix_config.json (chmod 600):
#   {"access_token": "...", "device_id": "..."}
# So kann das Skript ersetzt/geteilt werden, ohne Tokens anzufassen.
_cfg = json.loads((BASE / "matrix_config.json").read_text())
ACCESS_TOKEN = _cfg["access_token"]
DEVICE_ID = _cfg["device_id"]
STORE_PATH = str(BASE / "crypto_store")
ATTACHMENT_DIR = BASE / "attachments"
PENDING_FILE = BASE / "pending_megolm.json"
PENDING_MAX_AGE_DAYS = 14

# Raum-ID -> Fach-Zuordnung
ROOM_TO_FACH = {
    "!ZoLrbHDvIvNyrUVYwt:190688.logineonrw-messenger.de": "Mathe LK",
    "!THJCfpLWPulsfbLkLW:190688.logineonrw-messenger.de": "Sport LK",
    "!riHOtrDxjeWBcwcixA:190688.logineonrw-messenger.de": "Physik GK",
    "!HKCxGjpqUVTQLnPFpL:190688.logineonrw-messenger.de": "Deutsch GK",
    "!OOWOCPtkLgxTPdkzwp:190688.logineonrw-messenger.de": "Englisch GK",
    "!WIsldBrpicajbrEtKF:190688.logineonrw-messenger.de": "Spanisch",
    "!yWGivqamHdVQTygEms:190688.logineonrw-messenger.de": "SoWi ZK",
    "!irFPPovAqVMVVPyhYl:190688.logineonrw-messenger.de": "Geschichte GK",
    "!tGueZcDVowMJyxYIqZ:190688.logineonrw-messenger.de": "Informatik",
    "!kTihMALPKxwtiNRNvx:190688.logineonrw-messenger.de": "Astronomie",
    "!ZWmPdSDlDcKzNKdhLB:190688.logineonrw-messenger.de": "Abi27 Orga",
    "!GKCIOdBqvkQRRQRcwx:190688.logineonrw-messenger.de": "Stufenchat",
}

SYNC_FILTER = {"room": {"timeline": {"limit": 100}}}
MAX_BACKFILL_PAGES = 5

PLAIN_FILES = (RoomMessageFile, RoomMessageImage, RoomMessageVideo, RoomMessageAudio)
ENCRYPTED_FILES = (RoomEncryptedFile, RoomEncryptedImage, RoomEncryptedVideo, RoomEncryptedAudio)


def log(*args):
    print(*args, file=sys.stderr)


TEXT_MAX_ZEICHEN = 5000  # pro Anhang, damit Claude nicht mit riesigen Skripten zugeschüttet wird


def text_auszug(pfad: Path) -> str | None:
    """Text aus PDF/Word/Textdateien lesen, damit Claude den Inhalt von
    Arbeitsblättern sieht. Bilder u.ä. werden (noch) nicht gelesen."""
    endung = pfad.suffix.lower()
    try:
        if endung == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((seite.extract_text() or "") for seite in PdfReader(str(pfad)).pages[:15])
        elif endung == ".docx":
            import docx  # Paket python-docx
            d = docx.Document(str(pfad))
            teile = [p.text for p in d.paragraphs]
            for tabelle in d.tables:
                for zeile in tabelle.rows:
                    teile.append(" | ".join(zelle.text.strip() for zelle in zeile.cells))
            text = "\n".join(teile)
        elif endung in (".txt", ".md", ".csv"):
            text = pfad.read_text(encoding="utf-8", errors="replace")
        else:
            return None
    except ImportError as e:
        log(f"Textauszug nicht möglich, Paket fehlt: {e}")
        return None
    except Exception as e:  # noqa: BLE001  (kaputte/verschlüsselte PDFs usw.)
        log(f"Textauszug fehlgeschlagen ({pfad.name}): {e!r}")
        return None
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if not text:
        return None
    return text[:TEXT_MAX_ZEICHEN] + (" …[gekürzt]" if len(text) > TEXT_MAX_ZEICHEN else "")


def safe_filename(name: str, ts: int) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name or "datei")).strip(" .") or "datei"
    # Zeitstempel davor, damit "image.png" o.ä. sich nicht gegenseitig überschreiben
    return f"{time.strftime('%Y%m%d-%H%M%S', time.localtime(ts / 1000))}_{name}"


def sender_name(client: AsyncClient, room_id: str, sender: str) -> str:
    room = client.rooms.get(room_id)
    try:
        return room.user_name(sender) if room else sender
    except Exception:  # noqa: BLE001
        return sender


def load_pending() -> list:
    try:
        return json.loads(PENDING_FILE.read_text())
    except (FileNotFoundError, json.JSONDecodeError):
        return []


def save_pending(items: list) -> None:
    PENDING_FILE.write_text(json.dumps(items, ensure_ascii=False, indent=1))


async def key_maintenance(client: AsyncClient) -> None:
    """Nur das Nötigste, damit der Pi Schlüssel EMPFANGEN kann – nichts, was
    Nachrichten an fremde Geräte schickt.

    - keys_upload: veröffentlicht die öffentlichen Schlüssel DIESES Geräts
      (nur nötig, falls noch nicht geschehen; betrifft niemand anderen).
    - keys_query: reine Abfrage, welche Geräte die Raum-Mitglieder haben.

    Bewusst NICHT: keys_claim, send_to_device_messages, request_room_key.
    """
    try:
        if client.should_upload_keys:
            await client.keys_upload()
            log("E2EE: Geräteschlüssel hochgeladen")
        if client.should_query_keys:
            await client.keys_query()
            log("E2EE: Geräteliste der Raum-Mitglieder abgefragt")
    except Exception as e:  # noqa: BLE001
        log(f"E2EE-Wartung fehlgeschlagen: {e!r}")


async def handle_event(client, room_id, fach, event, results, pending):
    ts = event.server_timestamp
    base = {
        "fach": fach,
        "sender": event.sender,
        "sender_name": sender_name(client, room_id, event.sender),
        "timestamp": ts,
        "event_id": event.event_id,
    }

    if isinstance(event, RoomMessageText):
        results.append({**base, "type": "text", "body": event.body})

    elif isinstance(event, ENCRYPTED_FILES + PLAIN_FILES):
        orig_name = event.source.get("content", {}).get("filename") or event.body
        filename = safe_filename(orig_name, ts)
        filepath = ATTACHMENT_DIR / filename
        resp = await client.download(mxc=event.url)
        if isinstance(resp, DownloadError):
            log(f"Download fehlgeschlagen ({fach}, {orig_name}): {resp.message}")
            results.append({**base, "type": "attachment_error", "filename": orig_name})
            return
        data = resp.body
        if isinstance(event, ENCRYPTED_FILES):
            try:
                data = decrypt_attachment(data, event.key["k"], event.hashes["sha256"], event.iv)
            except Exception as e:  # noqa: BLE001
                log(f"Entschlüsseln fehlgeschlagen ({fach}, {orig_name}): {e!r}")
                results.append({**base, "type": "attachment_error", "filename": orig_name})
                return
        filepath.write_bytes(data)
        results.append({**base, "type": "attachment", "filename": filename,
                        "original_name": orig_name, "local_path": str(filepath),
                        "text_auszug": text_auszug(filepath)})

    elif isinstance(event, MegolmEvent):
        # Schlüssel fehlt (noch) -> merken und später erneut versuchen
        # Bewusst KEINE Schlüssel-Anfrage – der Pi liest nur.
        if not any(p["event_id"] == event.event_id for p in pending):
            pending.append({"room_id": room_id, "event_id": event.event_id,
                            "session_id": event.session_id, "fach": fach,
                            "timestamp": ts})


async def backfill(client, room_id, prev_batch, since_token) -> list:
    """Falls die Timeline trotz Limit abgeschnitten ist: ältere neue Events nachladen."""
    events, start = [], prev_batch
    for _ in range(MAX_BACKFILL_PAGES):
        resp = await client.room_messages(room_id, start=start, end=since_token, limit=100)
        if isinstance(resp, RoomMessagesError) or not resp.chunk:
            break
        events.extend(resp.chunk)
        if not resp.end or resp.end == start:
            break
        start = resp.end
    events.reverse()  # room_messages liefert rückwärts
    log(f"  Nachgeladen für {ROOM_TO_FACH[room_id]}: {len(events)} Events")
    return events


async def retry_pending(client, results, pending) -> list:
    still = []
    cutoff = (time.time() - PENDING_MAX_AGE_DAYS * 86400) * 1000
    for p in pending:
        if p.get("timestamp", 0) < cutoff:
            continue  # zu alt, aufgeben
        resp = await client.room_get_event(p["room_id"], p["event_id"])
        if isinstance(resp, RoomGetEventError):
            still.append(p)
            continue
        event = resp.event
        if isinstance(event, MegolmEvent):
            try:
                event = client.decrypt_event(event)
            except Exception:  # noqa: BLE001
                still.append(p)
                continue
        log(f"Nachträglich entschlüsselt: {p['fach']} {p['event_id']}")
        await handle_event(client, p["room_id"], p["fach"], event, results, [])
    return still


async def main():
    ATTACHMENT_DIR.mkdir(parents=True, exist_ok=True)
    client = AsyncClient(HOMESERVER, USER_ID, store_path=STORE_PATH,
                         config=AsyncClientConfig(store_sync_tokens=False, encryption_enabled=True))
    client.restore_login(USER_ID, DEVICE_ID, ACCESS_TOKEN)  # lädt auch den Crypto-Store

    since_token = load_token()
    resp = await client.sync(timeout=10000, since=since_token, sync_filter=SYNC_FILTER,
                             full_state=since_token is None,
                             set_presence="offline")  # nicht als "online" erscheinen
    if isinstance(resp, SyncError):
        log(f"Sync fehlgeschlagen: {resp.message}")
        await client.close()
        sys.exit(1)

    await key_maintenance(client)

    results, pending = [], load_pending()

    log("DEBUG: Räume mit Events:")
    for room_id, room_info in resp.rooms.join.items():
        if room_id not in ROOM_TO_FACH:
            continue
        fach = ROOM_TO_FACH[room_id]
        events = list(room_info.timeline.events)
        typen = {}
        for e in events:
            typen[type(e).__name__] = typen.get(type(e).__name__, 0) + 1
        log(f"  {fach}: {typen}" + (" (abgeschnitten)" if room_info.timeline.limited else ""))
        if room_info.timeline.limited and since_token:
            events = await backfill(client, room_id, room_info.timeline.prev_batch, since_token) + events
        seen = set()
        for event in events:
            if event.event_id in seen:
                continue
            seen.add(event.event_id)
            await handle_event(client, room_id, fach, event, results, pending)

    pending = await retry_pending(client, results, pending)
    save_pending(pending)
    if pending:
        log(f"Noch nicht entschlüsselbar: {len(pending)} Nachrichten (siehe pending_megolm.json)")

    save_token(resp.next_batch)
    await client.close()

    results.sort(key=lambda r: r.get("timestamp") or 0)
    print(json.dumps(results, ensure_ascii=False))


if __name__ == "__main__":
    asyncio.run(main())
